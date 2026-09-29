from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .perfiles import PerfilFormato, normalizar_clave


ALIASES_ENCABEZADO = {
    "seccion": ["SECCION", "SECCIÓN", "SECCI_N", "SECC", "SECCION_ELECTORAL", "NUM_SECCION"],
    "lista": ["LISTA_NOMINAL", "LISTA_NOMINAL_CASILLA", "LISTADO_NOMINAL", "LN"],
    "votos": ["TOTAL_VOTOS", "VOTOS_EMITIDOS", "TOTAL_VOTACION", "VOTACION_TOTAL", "VOTOS_TOTALES"],
    "nulos": ["NULOS", "NUM_VOTOS_NULOS", "VOTOS_NULOS", "VOTO_NULO"],
    "municipio": ["MUNICIPIO", "MUNICIPIO_LOCAL", "NOMBRE_MUNICIPIO", "NOM_MUNICIPIO"],
}


@dataclass(frozen=True)
class MetadataPerfil:
    fila_encabezado: int
    hoja: str
    filas_leidas: int
    encoding: str | None = None


def leer_tabla_perfil(
    ruta: Path,
    perfil: PerfilFormato | None = None,
    hoja: str | None = None,
    fila_encabezado: int | None = None,
) -> tuple[pd.DataFrame, MetadataPerfil]:
    """Lee una tabla electoral aunque tenga portada, varias hojas o CSV con preámbulo.

    ``perfil`` sólo aporta aliases y señales adicionales para favorecer una hoja; la
    detección de SECCION y de los encabezados básicos no depende del estado ni año.
    ``hoja`` y ``fila_encabezado`` (1-based) permiten forzar la lectura cuando la
    detección automática no elige la tabla correcta.
    """
    extension = ruta.suffix.lower()
    if extension == ".csv":
        return _leer_csv(ruta, perfil, fila_encabezado)
    if extension in {".xlsx", ".xlsm"}:
        return _leer_excel(ruta, perfil, hoja, fila_encabezado)
    raise ValueError("El archivo origen debe ser .xlsx, .xlsm o .csv.")


def listar_hojas(ruta: Path) -> list[str]:
    if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
        return []
    wb = load_workbook(ruta, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _leer_csv(
    ruta: Path, perfil: PerfilFormato | None, fila_forzada: int | None = None
) -> tuple[pd.DataFrame, MetadataPerfil]:
    ultimo_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin1"):
        try:
            bruto = _leer_csv_sin_encabezado(ruta, encoding)
            fila = fila_forzada or _detectar_encabezado_dataframe(bruto, "CSV", perfil)
            df = _dataframe_desde_fila_encabezado(bruto, fila)
            df = _limpiar_dataframe(df, perfil)
            return df, MetadataPerfil(fila_encabezado=fila, hoja="CSV", filas_leidas=len(df), encoding=encoding)
        except (UnicodeDecodeError, csv.Error, ValueError) as exc:
            ultimo_error = exc
    raise ValueError(f"No se pudo leer una tabla en el CSV: {ultimo_error}")


def _leer_csv_sin_encabezado(ruta: Path, encoding: str) -> pd.DataFrame:
    """Lee CSV con preámbulos de longitud variable sin asumir su primera fila.

    Pandas infiere el número de columnas a partir de la primera línea; eso falla en
    exportaciones que anteponen títulos de una sola celda. ``csv.reader`` conserva
    todas las filas y DataFrame rellena las celdas faltantes de forma segura.
    """
    with ruta.open("r", encoding=encoding, newline="") as archivo:
        muestra = archivo.read(65_536)
        archivo.seek(0)
        delimitadores = ";,\t|"
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=delimitadores)
            lector = csv.reader(archivo, dialecto)
        except csv.Error:
            delimitador = max(delimitadores, key=muestra.count)
            lector = csv.reader(archivo, delimiter=delimitador)
        filas = list(lector)
    if not filas:
        raise ValueError("El CSV está vacío.")
    return pd.DataFrame(filas)


def _leer_excel(
    ruta: Path,
    perfil: PerfilFormato | None,
    hoja_forzada: str | None = None,
    fila_forzada: int | None = None,
) -> tuple[pd.DataFrame, MetadataPerfil]:
    if fila_forzada and hoja_forzada:
        fila, hoja = fila_forzada, hoja_forzada
    else:
        fila, hoja = _detectar_tabla_excel(ruta, perfil, hoja_forzada)
        fila = fila_forzada or fila
    df = pd.read_excel(ruta, sheet_name=hoja, header=fila - 1)
    df = _limpiar_dataframe(df, perfil)
    return df, MetadataPerfil(fila_encabezado=fila, hoja=hoja, filas_leidas=len(df))


def _detectar_tabla_excel(
    ruta: Path, perfil: PerfilFormato | None, hoja: str | None = None
) -> tuple[int, str]:
    wb = load_workbook(ruta, read_only=True, data_only=True)
    mejor: tuple[int, int, str] | None = None
    respaldo: tuple[int, int, str] | None = None
    try:
        hojas = [wb[hoja]] if hoja else wb.worksheets
        for ws in hojas:
            for numero_fila, row in enumerate(ws.iter_rows(values_only=True), start=1):
                valores = [normalizar_clave(valor) for valor in row if valor is not None and str(valor).strip()]
                if numero_fila <= FILAS_RESPALDO:
                    puntaje_respaldo = _puntuar_encabezado_sin_seccion(list(row))
                    if puntaje_respaldo and (respaldo is None or puntaje_respaldo > respaldo[0]):
                        respaldo = (puntaje_respaldo, numero_fila, ws.title)
                puntaje = _puntuar_encabezado(valores, ws.title, perfil)
                if puntaje <= 0:
                    continue
                candidato = (puntaje, numero_fila, ws.title)
                if mejor is None or candidato[0] > mejor[0] or (candidato[0] == mejor[0] and candidato[1] < mejor[1]):
                    mejor = candidato
                # Un encabezado con SECCION y un conjunto amplio de columnas no
                # mejora materialmente al recorrer miles de filas de datos.
                if puntaje >= 50:
                    break
    finally:
        wb.close()

    if mejor is None:
        # Sin encabezado de SECCION reconocible: se toma la fila con más textos
        # para que la persona usuaria elija manualmente las columnas clave.
        mejor = respaldo
    if mejor is None:
        raise ValueError("No se detectó una fila de encabezados en ninguna hoja.")
    _, fila, hoja_detectada = mejor
    return fila, hoja_detectada


def _detectar_encabezado_dataframe(bruto: pd.DataFrame, hoja: str, perfil: PerfilFormato | None) -> int:
    mejor: tuple[int, int] | None = None
    respaldo: tuple[int, int] | None = None
    for indice, row in bruto.iterrows():
        if indice < FILAS_RESPALDO:
            puntaje_respaldo = _puntuar_encabezado_sin_seccion(row.tolist())
            if puntaje_respaldo and (respaldo is None or puntaje_respaldo > respaldo[0]):
                respaldo = (puntaje_respaldo, indice + 1)
        valores = [normalizar_clave(valor) for valor in row.tolist() if pd.notna(valor) and str(valor).strip()]
        puntaje = _puntuar_encabezado(valores, hoja, perfil)
        if puntaje <= 0:
            continue
        candidato = (puntaje, indice + 1)
        if mejor is None or candidato[0] > mejor[0] or (candidato[0] == mejor[0] and candidato[1] < mejor[1]):
            mejor = candidato
        if puntaje >= 50:
            break
    if mejor is None:
        mejor = respaldo
    if mejor is None:
        raise ValueError("No se detectó una fila de encabezados.")
    return mejor[1]


FILAS_RESPALDO = 100


def _puntuar_encabezado_sin_seccion(valores: list[object]) -> int:
    """Cuenta celdas de texto no numérico; un encabezado típico tiene varias."""
    textos = 0
    for valor in valores:
        if valor is None or (isinstance(valor, float) and pd.isna(valor)):
            continue
        texto = str(valor).strip()
        if not texto:
            continue
        try:
            float(texto.replace(",", ""))
        except ValueError:
            textos += 1
    return textos if textos >= 3 else 0


def _dataframe_desde_fila_encabezado(bruto: pd.DataFrame, fila: int) -> pd.DataFrame:
    encabezados = [str(valor).strip() if pd.notna(valor) else "" for valor in bruto.iloc[fila - 1].tolist()]
    datos = bruto.iloc[fila:].copy()
    datos.columns = encabezados
    return datos


def _puntuar_encabezado(valores: list[str], hoja: str, perfil: PerfilFormato | None) -> int:
    conjunto = set(valores)
    aliases = {clave: list(valores_alias) for clave, valores_alias in ALIASES_ENCABEZADO.items()}
    if perfil is not None:
        for clave, valores_alias in perfil.aliases_columnas.items():
            aliases.setdefault(clave, [])
            aliases[clave].extend(valores_alias)

    aliases_normalizados = {
        clave: {normalizar_clave(alias) for alias in valores_alias}
        for clave, valores_alias in aliases.items()
    }
    if not conjunto & aliases_normalizados["seccion"]:
        return 0

    # SECCION es indispensable. Los demás campos distinguen encabezados reales de
    # títulos, notas o tablas auxiliares que sólo mencionan la palabra "sección".
    puntaje = 30
    # Un encabezado tabular suele contener varias columnas, aun cuando el estado
    # use nombres de métricas distintos a los aliases conocidos.
    puntaje += min(len(conjunto), 20)
    for clave in ("lista", "votos", "nulos", "municipio"):
        if conjunto & aliases_normalizados.get(clave, set()):
            puntaje += 15

    if perfil is not None:
        for partido in perfil.partidos_salida:
            aliases_partido = perfil.aliases_columnas.get(partido, [partido])
            if conjunto & {normalizar_clave(alias) for alias in aliases_partido}:
                puntaje += 2

    hoja_norm = normalizar_clave(hoja)
    if hoja_norm in {"SECCION", "SECCIONES", "DATOS"} or hoja_norm.endswith("_SECCION"):
        puntaje += 5
    if "CASILLA" in hoja_norm:
        puntaje += 2
    return puntaje


def _limpiar_dataframe(df: pd.DataFrame, perfil: PerfilFormato | None) -> pd.DataFrame:
    df = df.replace(r"^\s*$", pd.NA, regex=True)
    df = df.dropna(axis=1, how="all").dropna(how="all").copy()
    df.columns = _encabezados_unicos(df.columns)

    aliases_seccion = (perfil.aliases_columnas.get("seccion", []) if perfil is not None else []) + ALIASES_ENCABEZADO["seccion"]
    seccion_col = _resolver_columna(df, aliases_seccion)
    if seccion_col is None:
        # Sin SECCION reconocible se devuelve la tabla completa; el formateador
        # exigirá que la persona usuaria elija la columna de sección.
        return df.reset_index(drop=True)
    if seccion_col != "SECCION" and "SECCION" not in df.columns:
        df = df.rename(columns={seccion_col: "SECCION"})
    df["SECCION"] = pd.to_numeric(df["SECCION"], errors="coerce")
    df = df.dropna(subset=["SECCION"]).copy()
    df["SECCION"] = df["SECCION"].astype(int)
    return df.reset_index(drop=True)


def _encabezados_unicos(columnas) -> list[str]:
    """Nombres no vacíos y únicos (``PAN``, ``PAN.1``) para evitar columnas ambiguas."""
    resultado: list[str] = []
    vistos: set[str] = set()
    for indice, col in enumerate(columnas, start=1):
        nombre = "" if col is None or (isinstance(col, float) and pd.isna(col)) else str(col).strip()
        if not nombre or nombre.upper().startswith("UNNAMED:"):
            nombre = f"COLUMNA_{indice}"
        base, n = nombre, 1
        while nombre in vistos:
            nombre = f"{base}.{n}"
            n += 1
        vistos.add(nombre)
        resultado.append(nombre)
    return resultado


def _resolver_columna(df: pd.DataFrame, aliases: list[str]) -> str | None:
    indice = {normalizar_clave(col): col for col in df.columns}
    for alias in aliases:
        col = indice.get(normalizar_clave(alias))
        if col is not None:
            return col
    return None
