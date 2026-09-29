from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.entidades import detectar_entidad
from src.excel_writer import escribir_xlsx
from src.lector_perfiles import leer_tabla_perfil
from src.mapeo_columnas import ColumnasClaveFaltantes, detectar_columnas_clave, detectar_partidos
from src.perfiles import crear_perfil_generico, detectar_tipo_eleccion
from src.sabana import formatear_sabana

ROOT = Path(__file__).resolve().parents[1]
REF_DIP_QRO_2024 = ROOT / "data" / "SE_DIP_LOCALES_QRO_2024.xlsx"
REF_AYUN_QRO_2021 = ROOT / "dataPrueba" / "base_electoral_formateada_ayuntamientos_2021 (1).xlsx"


def _encabezados_referencia(ruta: Path) -> list[str]:
    libro = load_workbook(ruta, read_only=True)
    try:
        return [c for c in next(libro["Formato"].iter_rows(max_row=1, values_only=True)) if c is not None]
    finally:
        libro.close()


def _partidos_de_referencia(encabezados: list[str]) -> list[str]:
    inicio = encabezados.index("3PP_MV") + 3
    fin = encabezados.index("CNR")
    return [h for h in encabezados[inicio:fin] if h != "PCN"]


def _encabezados_generados(nombre_archivo: str, partidos: list[str]) -> list[str]:
    datos = {"ID_ESTADO": [22], "NOMBRE_ESTADO": ["QUERETARO"], "MUNICIPIO": ["AMEALCO"], "SECCION": [1],
             "LISTA_NOMINAL": [100], "TOTAL_VOTOS": [10 * len(partidos) + 3], "NULOS": [2], "NO_REGISTRADOS": [1]}
    datos.update({p: [10] for p in partidos})
    resultado = formatear_sabana(pd.DataFrame(datos), crear_perfil_generico(nombre_archivo), partidos=partidos)
    hoja = load_workbook(BytesIO(escribir_xlsx(resultado.df_base, resultado.config)))["Formato"]
    return [c.value for c in hoja[1]]


@pytest.mark.parametrize(
    ("referencia", "nombre_archivo"),
    [(REF_DIP_QRO_2024, "DIP_LOC_QRO_2024.csv"), (REF_AYUN_QRO_2021, "AYUN_QRO_2021.csv")],
)
def test_encabezados_identicos_a_las_salidas_de_referencia(referencia: Path, nombre_archivo: str):
    if not referencia.exists():
        pytest.skip(f"No existe {referencia}")
    esperados = _encabezados_referencia(referencia)
    assert _encabezados_generados(nombre_archivo, _partidos_de_referencia(esperados)) == esperados


def test_detecta_tipo_de_eleccion_y_entidad_por_nombre():
    assert detectar_tipo_eleccion("GTO_GUB_2024.csv") == "gubernatura"
    assert detectar_tipo_eleccion("GTO_GOB_2018.xlsx") == "gubernatura"
    assert detectar_tipo_eleccion("GTO_DIP_LOC_2024.csv") == "diputaciones"
    assert detectar_tipo_eleccion("Ayuntamiento 2021.xlsx") == "ayuntamientos"
    assert detectar_tipo_eleccion("QRO_AYUN_RESULTADOS_2024.csv") == "ayuntamientos"
    assert detectar_entidad("GTO_GUB_2024.csv") == (11, "GUANAJUATO")
    assert detectar_entidad("SE_DIP_LOCALES_QRO_2024.xlsx") == (22, "QUERETARO")
    assert detectar_entidad("resultados_baja_california_sur_2021.csv") == (3, "BAJA CALIFORNIA SUR")
    assert detectar_entidad("2021_Gubernatura.xlsx") is None
    assert crear_perfil_generico("resultados.csv").anio == ""


def test_columnas_duplicadas_de_cnr_y_nulos_no_se_cuentan_doble():
    # Formato de cómputos de Guanajuato: NoReg/Nulos y NUM_VOTOS_CAN_NREG/NUM_VOTOS_NULOS repiten el dato,
    # y cada partido trae su porcentaje P_<PARTIDO>.
    tabla = pd.DataFrame({
        "ID_ESTADO": [11, 11], "NOMBRE_ESTADO": ["GUANAJUATO"] * 2, "ID_DISTRITO_LOCAL": ["01", "01"],
        "ID_UBICACION": [47, 47], "UBICACION": ["ABASOLO"] * 2, "SECCION": ["0752", "0752"],
        "TIPO_CASILLA": ["B", "C"], "ID_CASILLA": ["01", "01"], "EXT_CONTIGUA": ["00", "00"],
        "PAN": [165, 167], "P_PAN": [48.8165, 52.5157], "PRI": [160, 144], "P_PRI": [47.3, 45.2],
        "CAND_IND_1": [0, 0], "P_CAND_IND_1": [0.0, 0.0], "NoReg": [7, 0], "P_NoReg": [2.1, 0.0],
        "Nulos": [6, 7], "P_Nulos": [1.7, 2.2], "NUM_VOTOS_VALIDOS": [325, 311], "NUM_VOTOS_CAN_NREG": [7, 0],
        "NUM_VOTOS_NULOS": [6, 7], "TOTAL_VOTOS": [338, 318], "LISTA_NOMINAL": [617, 617],
        "PARTICIPACION_CONTABILIZADA": [54.78, 51.53], "COTEJADA": [1, 1], "IDCASILLA": [1888, 1889],
    })
    resultado = formatear_sabana(tabla, crear_perfil_generico("GTO_GUB_2024.csv"))
    fila = resultado.df_base.iloc[0]

    assert resultado.partidos == ["PAN", "PRI", "CAND_IND_1"]
    assert fila["SECCION"] == 752
    assert fila["DL"] == 1
    assert fila["CNR"] == 7
    assert fila["NULOS"] == 13
    assert fila["PAN"] + fila["PRI"] + fila["CNR"] + fila["NULOS"] == fila["VOTOS_EMITIDOS"] == 656
    assert not any("VALIDACION" in aviso for aviso in resultado.advertencias)


def test_independientes_por_iniciales_y_no_registradas():
    # Formato de Querétaro 2024: candidaturas independientes con iniciales y NO_REGISTRADAS.
    tabla = pd.DataFrame({
        "ID_ESTADO": ["22"], "ESTADO": ["QUERETARO"], "ID_MUNICIPIO": ["001"], "MUNICIPIO": ["Amealco"],
        "SECCION": [1], "ID_CASILLA": ["01"], "TIPO_CASILLA": ["B"], "UBICACION_CASILLA": ["URBANA"],
        "PAN": ["51"], "MC": ["190"], "EMC": ["4"], "JBLL": ["0"], "PVEM-MORENA-PT": ["5"],
        "NO_REGISTRADAS": ["1"], "NULOS": ["13"], "TOTAL_VOTOS": ["264"], "LISTA_NOMINAL": ["536"],
        "OBSERVACIONES": [""],
    })
    resultado = formatear_sabana(tabla, crear_perfil_generico("QRO_AYUN_RESULTADOS_2024.csv"))

    assert resultado.partidos == ["PAN", "MC", "EMC", "JBLL", "PVEM-MORENA-PT"]
    assert resultado.mapeo["CNR"] == "NO_REGISTRADAS"
    assert resultado.df_base.loc[0, "MUNICIPIO"] == "Amealco"
    assert "MUNICIPIO" in resultado.config.encabezados_visibles


def test_prefijo_p_sin_columna_homologa_son_votos_de_coalicion():
    tabla = pd.DataFrame({
        "SECCION": [1], "LISTA_NOMINAL": [100], "VOTOS_EMITIDOS": [60], "VOTOS_NULOS": [5],
        "MORENA": [30], "P_MORENA": [50.0], "PT": [20], "P_MORENA-PT": [5],
    })
    assert detectar_partidos(tabla, detectar_columnas_clave(tabla)) == ["MORENA", "PT", "P_MORENA-PT"]


def test_columnas_clave_no_detectadas_se_seleccionan_manualmente():
    tabla = pd.DataFrame({
        "CLAVE_GEO": [5, 5, 6], "ELECTORES": [300, 200, 400], "BOLETAS_EN_URNA": [150, 100, 200],
        "INVALIDOS": [5, 5, 10], "LUGAR_MUN": ["LEON", "LEON", "LEON"],
        "PAN": [100, 60, 120], "MORENA": [45, 35, 70],
    })
    perfil = crear_perfil_generico("GTO_AYUN_2024.csv")

    with pytest.raises(ColumnasClaveFaltantes) as error:
        formatear_sabana(tabla, perfil)
    assert set(error.value.faltantes) == {"seccion", "lista", "votos", "nulos"}

    mapeo = {"seccion": "CLAVE_GEO", "lista": "ELECTORES", "votos": "BOLETAS_EN_URNA",
             "nulos": "INVALIDOS", "municipio": "LUGAR_MUN"}
    resultado = formatear_sabana(tabla, perfil, mapeo=mapeo, partidos=["PAN", "MORENA"])
    fila = resultado.df_base.iloc[0]

    assert list(resultado.df_base["SECCION"]) == [5, 6]
    assert fila["LISTA_NOMINAL"] == 500
    assert fila["VOTOS_EMITIDOS"] == 250
    assert fila["NULOS"] == 10
    assert fila["MUNICIPIO"] == "LEON"
    # La entidad sale del nombre del archivo porque no hay columna.
    assert (fila["CVE_ENTIDAD"], fila["ENTIDAD"]) == (11, "GUANAJUATO")


def test_lector_devuelve_tabla_aunque_no_haya_encabezado_de_seccion(tmp_path: Path):
    ruta = tmp_path / "resultados_2024.csv"
    ruta.write_text(
        "Resultados\nCLAVE_GEO,ELECTORES,BOLETAS_EN_URNA,INVALIDOS,PAN\n5,300,150,5,100\n6,400,200,10,120\n",
        encoding="utf-8",
    )
    tabla, meta = leer_tabla_perfil(ruta)
    assert meta.fila_encabezado == 2
    assert list(tabla.columns) == ["CLAVE_GEO", "ELECTORES", "BOLETAS_EN_URNA", "INVALIDOS", "PAN"]


def test_fila_de_encabezado_forzada(tmp_path: Path):
    ruta = tmp_path / "x_2024.csv"
    ruta.write_text("A,B,C\nSECCION,LISTA_NOMINAL,PAN\n1,10,5\n", encoding="utf-8")
    tabla, meta = leer_tabla_perfil(ruta, fila_encabezado=2)
    assert meta.fila_encabezado == 2
    assert tabla.loc[0, "PAN"] == "5"


def test_formulas_de_validacion_incluyen_cnr_y_nulos():
    tabla = pd.DataFrame({"SECCION": [1], "LISTA_NOMINAL": [100], "TOTAL_VOTOS": [60], "NULOS": [5],
                          "CNR": [1], "PAN": [30], "PRI": [24]})
    resultado = formatear_sabana(tabla, crear_perfil_generico("DIP_2024.csv"))
    hoja = load_workbook(BytesIO(escribir_xlsx(resultado.df_base, resultado.config)))["Formato"]
    encabezados = [c.value for c in hoja[1]]

    assert encabezados[-6:] == ["CNR", "PCN", "NULOS", "PCN", "TOT_VOTOS", "VALIDACION"]
    tot = hoja.cell(2, encabezados.index("TOT_VOTOS") + 1).value
    for header in ("PAN", "PRI", "CNR", "NULOS"):
        letra = hoja.cell(1, encabezados.index(header) + 1).column_letter
        assert f"{letra}2" in tot


@pytest.mark.parametrize(
    ("ruta", "secciones"),
    [
        (ROOT / "data" / "2021_Gubernatura.xlsx", 892),
        (ROOT / "data" / "Ayuntamiento 2021.xlsx", 891),
        (ROOT / "data" / "QRO_AYUN_RESULTADOS_2024.csv", 953),
        (ROOT / "dataPruebaGua" / "09062024_1200_COMPUTOS_GUB_GTO" / "GTO_GUB_2024.csv", 3255),
        (ROOT / "dataPruebaGua" / "GTO_AYU_2018.xlsx", 3137),
    ],
)
def test_archivos_reales_validan_al_100(ruta: Path, secciones: int):
    if not ruta.exists():
        pytest.skip(f"No existe {ruta}")
    perfil = crear_perfil_generico(ruta.name)
    tabla, _ = leer_tabla_perfil(ruta, perfil)
    resultado = formatear_sabana(tabla, perfil)
    df = resultado.df_base

    assert len(df) == secciones
    total = df[resultado.config.partidos + ["NULOS"]].sum(axis=1)
    assert (total == df["VOTOS_EMITIDOS"]).all()
    assert df["CVE_ENTIDAD"].ne("").all()
