from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.excel_writer import escribir_xlsx
from src.formateador_simple import FormateadorSimple
from src.perfiles import crear_perfil_generico


PREFIJO_SABANA_2024 = [
    "#", "CVE_ENTIDAD", "ENTIDAD", "DF", "DL", "SECCION", "LISTA_NOMINAL", "VOTOS_EMITIDOS",
    "PARTICIPACION", "ABSTENCION", "1ER_LUGAR", "1ERO_VOTOS", "PCN", "DIF_VOTOS_2DO",
    "DIF_PCN_2DO", "2DO_LUGAR", "2DO_VOTOS", "PCN", "DIF_VOTOS_3RO", "DIF_PCN_3RO",
    "3ER_LUGAR", "3RO_VOTOS", "PCN", "1PP_MV", "VOTOS", "PCN", "DIF_2DO", "PCN",
    "2PP_MV", "VOTOS", "PCN", "DIF_3RO", "PCN", "3PP_MV", "VOTOS", "PCN",
]


def test_generico_genera_sabana_2024_y_clasifica_metricas_indispensables():
    perfil = crear_perfil_generico("GTO_AYU_2018.xlsx")
    tabla = pd.DataFrame(
        {
            "ID_ESTADO": [11, 11],
            "NOMBRE_ESTADO": ["GUANAJUATO", "GUANAJUATO"],
            "DF": [5, 5],
            "DL": [7, 7],
            "SECCION": [101, 101],
            "LISTA_NOMINAL": [1000, 200],
            "TOTAL": [700, 150],
            "NULO": [12, 3],
            "NOREG": [4, 1],
            "PAN": [350, 70],
            "PRI": [334, 76],
        }
    )

    resultado = FormateadorSimple(perfil).formatear(tabla)
    fila = resultado.df_base.iloc[0]

    assert resultado.config.encabezados_visibles[: len(PREFIJO_SABANA_2024)] == PREFIJO_SABANA_2024
    assert "CU_MUNICIPIO" not in resultado.config.encabezados_visibles
    assert "MUNICIPIO" not in resultado.config.encabezados_visibles
    assert resultado.config.partidos == ["PAN", "PRI", "CNR"]
    assert fila["VOTOS_EMITIDOS"] == 850
    assert fila["NULOS"] == 15
    assert fila["CNR"] == 5
    assert fila["PAN"] == 420
    assert fila["PRI"] == 410

    libro = load_workbook(BytesIO(escribir_xlsx(resultado.df_base, resultado.config)), data_only=False)
    hoja = libro["Formato"]
    encabezados = [celda.value for celda in hoja[1]]
    assert encabezados == resultado.config.encabezados_visibles
    assert hoja["I2"].value == "=IFERROR(H2/G2,0)"
    assert hoja.cell(2, encabezados.index("TOT_VOTOS") + 1).value.startswith("=")
    assert hoja.cell(2, encabezados.index("VALIDACION") + 1).value.startswith("=")



def test_generico_excluye_porcentajes_pcn_y_campos_operativos_de_guanajuato():
    perfil = crear_perfil_generico("GTO_GUB_2024.xlsx")
    tabla = pd.DataFrame(
        {
            "ID_ENTIDAD": [11, 11],
            "SECCION": [101, 101],
            "LN": [1000, 200],
            "TOTAL": [700, 150],
            "NULOS": [12, 3],
            "NOREG": [4, 1],
            "PAN": [350, 70],
            "P_PAN": [50.0, 46.67],
            "PRI": [334, 76],
            "P_PRI": [47.71, 50.67],
            "PARTICIPACION_CONTABILIZADA": [0.70, 0.75],
            "PCN": [0.01, 0.02],
            "COTEJADA": [1, 1],
            "RECONTADA": [0, 1],
            "IDCASILLA": [6, 187],
            "CONTABILIZADA": [1, 1],
        }
    )

    resultado = FormateadorSimple(perfil).formatear(tabla)
    fila = resultado.df_base.iloc[0]

    assert resultado.config.partidos == ["PAN", "PRI", "CNR"]
    assert list(resultado.df_base.columns) == [
        "#", "CVE_ENTIDAD", "ENTIDAD", "CU_MUNICIPIO", "MUNICIPIO", "DF", "DL", "SECCION",
        "LISTA_NOMINAL", "VOTOS_EMITIDOS", "PAN", "PRI", "CNR", "NULOS",
    ]
    assert fila["PAN"] == 420
    assert fila["PRI"] == 410
    assert fila["CNR"] == 5
    assert fila["NULOS"] == 15

    libro = load_workbook(BytesIO(escribir_xlsx(resultado.df_base, resultado.config)), data_only=False)
    hoja = libro["Formato"]
    encabezados = [celda.value for celda in hoja[1]]
    total_columna = encabezados.index("TOT_VOTOS") + 1
    validacion_columna = encabezados.index("VALIDACION") + 1
    assert hoja.cell(2, encabezados.index("PCN", resultado.config.indice_inicio_partidos) + 1).value == "=IFERROR(AK2/H2,0)"
    assert hoja.cell(2, validacion_columna).value == f"=IFERROR({get_column_letter(total_columna)}2/H2,0)"
